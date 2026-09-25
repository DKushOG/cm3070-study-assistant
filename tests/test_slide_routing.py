"""Tests for the per-slide extraction routing rule.

Part of the per-slide routing group. The rule decides, before any model call,
whether a slide is worth sending to the vision model. These tests fix the
decision boundaries and the bias the rule is built on, so neither changes by
accident.

No test here calls a model or reads a real deck. The picture-reading tests
build a small deck of their own in a temporary folder.
"""
import os
import struct
import tempfile
import unittest
import zlib

from modules import slide_routing


class ThresholdTests(unittest.TestCase):
    """The thresholds in the code match the ones that were measured.

        Written because a threshold is a number in a file, and a number in a
        file drifts. These values came from the instrumented corpus, 81
        PowerPoint slides for picture area and 111 slides for text yield, so
        changing either should take new evidence rather than a tidy-up.

    """

    def test_picture_area_threshold_is_the_measured_one(self):
        self.assertAlmostEqual(slide_routing.PICTURE_AREA_THRESHOLD, 0.10)

    def test_text_fallback_threshold_is_the_measured_one(self):
        self.assertEqual(slide_routing.TEXT_WORDS_THRESHOLD, 60)


class PictureAreaRuleTests(unittest.TestCase):
    """A picture-heavy slide goes to the model and a text slide does not."""

    def test_a_picture_heavy_slide_goes_to_the_vision_model(self):
        route, reason = slide_routing.decide(picture_area_ratio=0.74)
        self.assertEqual(route, slide_routing.ROUTE_VISION)
        self.assertIn("74%", reason)

    def test_a_text_slide_is_read_from_the_text_layer(self):
        route, _reason = slide_routing.decide(picture_area_ratio=0.02)
        self.assertEqual(route, slide_routing.ROUTE_TEXT)

    def test_the_boundary_is_inclusive(self):
        # At exactly the threshold the slide is sent. Recall is preferred at
        # the margin because a slide wrongly skipped loses its diagram
        # permanently, while a slide wrongly sent costs a few seconds.
        route, _reason = slide_routing.decide(
            picture_area_ratio=slide_routing.PICTURE_AREA_THRESHOLD)
        self.assertEqual(route, slide_routing.ROUTE_VISION)

    def test_picture_area_wins_over_text_yield_when_both_are_known(self):
        """A dense slide that is mostly picture still goes to vision.

        Picture area correlates with the outcome at +0.674 and text yield at
        -0.485, so where they disagree the stronger predictor decides.
        """
        route, _reason = slide_routing.decide(picture_area_ratio=0.60,
                                              text_layer_words=190)
        self.assertEqual(route, slide_routing.ROUTE_VISION)


class TextYieldFallbackTests(unittest.TestCase):
    """For a PDF page, the word count of the text layer decides."""

    def test_a_sparse_pdf_page_goes_to_the_vision_model(self):
        route, reason = slide_routing.decide(text_layer_words=12)
        self.assertEqual(route, slide_routing.ROUTE_VISION)
        self.assertIn("12 words", reason)

    def test_a_wordy_pdf_page_is_read_from_the_text_layer(self):
        route, _reason = slide_routing.decide(text_layer_words=190)
        self.assertEqual(route, slide_routing.ROUTE_TEXT)

    def test_the_fallback_boundary(self):
        below, _ = slide_routing.decide(
            text_layer_words=slide_routing.TEXT_WORDS_THRESHOLD - 1)
        at, _ = slide_routing.decide(
            text_layer_words=slide_routing.TEXT_WORDS_THRESHOLD)
        self.assertEqual(below, slide_routing.ROUTE_VISION)
        self.assertEqual(at, slide_routing.ROUTE_TEXT)


class NoFeaturesTests(unittest.TestCase):
    """With nothing to go on, the slide is sent to the model."""

    def test_with_nothing_to_go_on_the_slide_is_sent(self):
        """Defaulting to vision is the deliberate choice, not an oversight.

        The two errors are not symmetric: skipping a slide that needed the
        model loses content silently and permanently, and calling one that
        did not costs about eight seconds.
        """
        route, reason = slide_routing.decide()
        self.assertEqual(route, slide_routing.ROUTE_VISION)
        self.assertIn("defaulting", reason)


class CustomThresholdTests(unittest.TestCase):
    """Both thresholds can be overridden, which is what an ablation needs."""

    def test_thresholds_can_be_overridden_for_an_ablation(self):
        route, _reason = slide_routing.decide(picture_area_ratio=0.30,
                                              picture_threshold=0.50)
        self.assertEqual(route, slide_routing.ROUTE_TEXT)


class PictureAreaReadingTests(unittest.TestCase):
    """A file that is not a readable pptx gives no ratios and no error."""

    def test_a_non_pptx_yields_no_ratios(self):
        # Which is what makes a PDF fall back to text yield rather than
        # silently routing everything one way.
        self.assertEqual(slide_routing.pptx_picture_area_ratios("deck.pdf"),
                         {})

    def test_a_missing_file_does_not_raise(self):
        self.assertEqual(
            slide_routing.pptx_picture_area_ratios("no_such_deck.pptx"), {})


def _tiny_png():
    """A valid one-pixel PNG, built by hand so no imaging library is needed."""
    def chunk(kind, data):
        body = kind + data
        return (struct.pack(">I", len(data)) + body
                + struct.pack(">I", zlib.crc32(body) & 0xFFFFFFFF))
    header = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
    pixels = zlib.compress(b"\x00\xff\xff\xff")
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", header)
            + chunk(b"IDAT", pixels) + chunk(b"IEND", b""))


try:
    import pptx  # noqa: F401
    HAVE_PPTX = True
except ImportError:
    HAVE_PPTX = False


@unittest.skipUnless(HAVE_PPTX, "python-pptx is not installed")
class PlaceholderPictureTests(unittest.TestCase):
    """A picture counts however it was placed on the slide.

        Written after a real lecture slide was scored as having no pictures.
        Its picture sat in a content placeholder, which the shape type misses.

    """

    def setUp(self):
        from pptx import Presentation
        self.folder = tempfile.TemporaryDirectory()
        self.image = os.path.join(self.folder.name, "pixel.png")
        with open(self.image, "wb") as handle:
            handle.write(_tiny_png())
        self.presentation = Presentation()

    def tearDown(self):
        self.folder.cleanup()

    def _save(self):
        path = os.path.join(self.folder.name, "deck.pptx")
        self.presentation.save(path)
        return path

    def _layout_with(self, placeholder_type):
        for layout in self.presentation.slide_layouts:
            for placeholder in layout.placeholders:
                if placeholder.placeholder_format.type == placeholder_type:
                    return layout, placeholder.placeholder_format.idx
        self.skipTest("default template has no such placeholder")

    def test_a_free_picture_counts(self):
        from pptx.util import Emu
        slide = self.presentation.slides.add_slide(
            self.presentation.slide_layouts[6])
        width = self.presentation.slide_width // 2
        height = self.presentation.slide_height // 2
        slide.shapes.add_picture(self.image, 0, 0, Emu(width), Emu(height))
        ratios = slide_routing.pptx_picture_area_ratios(self._save())
        self.assertAlmostEqual(ratios[1], 0.25, places=2)

    def test_a_picture_in_a_placeholder_counts(self):
        from pptx.enum.shapes import PP_PLACEHOLDER
        layout, idx = self._layout_with(PP_PLACEHOLDER.PICTURE)
        slide = self.presentation.slides.add_slide(layout)
        inserted = slide.placeholders[idx].insert_picture(self.image)
        self.assertTrue(slide_routing.is_picture_shape(inserted))
        expected = (float(inserted.width) * float(inserted.height)) / (
            float(self.presentation.slide_width)
            * float(self.presentation.slide_height))
        ratios = slide_routing.pptx_picture_area_ratios(self._save())
        self.assertGreater(ratios[1], 0.0)
        self.assertAlmostEqual(ratios[1], min(expected, 1.0), places=3)

    def test_a_text_placeholder_is_not_a_picture(self):
        # Every text slide has placeholders, so counting all of them would
        # send every slide to the vision model.
        slide = self.presentation.slides.add_slide(
            self.presentation.slide_layouts[1])
        for placeholder in slide.placeholders:
            placeholder.text = "Bayes rule and conditional probability"
            self.assertFalse(slide_routing.is_picture_shape(placeholder))
        ratios = slide_routing.pptx_picture_area_ratios(self._save())
        self.assertEqual(ratios[1], 0.0)


class SummaryTests(unittest.TestCase):
    """A routed run describes what it avoided, and an empty report does not."""

    def test_a_routed_run_describes_what_it_avoided(self):
        summary = slide_routing.summarise(
            {"pages": 100, "vision_pages": 57})
        self.assertIn("57 of 100", summary)
        self.assertIn("43%", summary)

    def test_an_empty_report_says_nothing(self):
        self.assertEqual(slide_routing.summarise({}), "")
        self.assertEqual(slide_routing.summarise(None), "")


if __name__ == "__main__":
    unittest.main()
