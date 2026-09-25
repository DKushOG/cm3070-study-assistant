"""Decide, per slide, whether to call the vision model.

Slides that are mostly pictures usually gain new content from a vision call,
while text slides mostly repeat what the text layer already holds.

The rule:
  PowerPoint, send the slide when pictures cover at least 10 per cent of it.
  PDF, where picture area cannot be read, send it when the text layer returns
  fewer than 60 words.

Both thresholds were set from measurements taken on the evaluation corpus.
Recall is favoured over precision, because a skipped slide loses its diagram
from the output while an unnecessary call only costs time.

evaluation/route_instrumentation.py imports this module, so both use the same
routing logic when deciding which slides need vision processing.
"""
from pathlib import Path

# Send to the vision model at or above this share of a slide. PowerPoint only.
PICTURE_AREA_THRESHOLD = 0.10

# PDF fallback: send the slide when the text layer gives fewer words than this.
TEXT_WORDS_THRESHOLD = 60

ROUTE_VISION = "vision"
ROUTE_TEXT = "text layer"

# Stamped on every page so the output shows which path produced it.
ROUTE_MARKER = "[extracted by: {route}]"


def is_picture_shape(shape):
    """True when a shape holds a picture.

    Counts a normal picture and a picture dropped into a content placeholder.
    PowerPoint stores the second kind as a placeholder, so checking the shape
    type alone misses it. Text, chart and table placeholders do not count.
    """
    try:
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except ImportError:
        return False
    try:
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            return True
    except NotImplementedError:
        # python-pptx raises this for some unrecognised shape kinds.
        pass
    if not getattr(shape, "is_placeholder", False):
        return False
    element = getattr(shape, "_element", None)
    tag = getattr(element, "tag", "")
    return isinstance(tag, str) and tag.endswith("}pic")


def pptx_picture_area_ratios(path):
    """Return {slide number: share of the slide covered by pictures}.

    Picture areas are read directly from the PowerPoint file without rendering
    the slides or calling a model. Returns {} for anything that is not a
    readable pptx, which is what makes the caller fall back to text yield.
    Overlapping pictures are summed and the total is capped at 1.0.
    """
    if Path(str(path)).suffix.lower() != ".pptx":
        return {}
    try:
        from pptx import Presentation
    except ImportError:
        return {}
    try:
        presentation = Presentation(str(path))
    except Exception:
        return {}
    slide_area = float(presentation.slide_width) * float(
        presentation.slide_height)
    if not slide_area:
        return {}
    ratios = {}
    for number, slide in enumerate(presentation.slides, start=1):
        covered = 0.0
        for shape in slide.shapes:
            if not is_picture_shape(shape):
                continue
            try:
                covered += float(shape.width) * float(shape.height)
            except (TypeError, ValueError):
                continue
        ratios[number] = min(covered / slide_area, 1.0)
    return ratios


def decide(picture_area_ratio=None, text_layer_words=None,
           picture_threshold=None, words_threshold=None):
    """Return (route, reason) for one slide.

    Use picture area for PowerPoint slides and text-layer word count for PDF
    pages. The reason is returned so a run can be checked afterwards rather
    than only counted.
    """
    if picture_threshold is None:
        picture_threshold = PICTURE_AREA_THRESHOLD
    if words_threshold is None:
        words_threshold = TEXT_WORDS_THRESHOLD

    if picture_area_ratio is not None:
        if picture_area_ratio >= picture_threshold:
            return ROUTE_VISION, (
                f"pictures cover {picture_area_ratio:.0%} of the slide, at or "
                f"above the {picture_threshold:.0%} threshold")
        return ROUTE_TEXT, (
            f"pictures cover only {picture_area_ratio:.0%} of the slide")

    if text_layer_words is not None:
        if text_layer_words < words_threshold:
            return ROUTE_VISION, (
                f"the text layer yields only {text_layer_words} words, below "
                f"the {words_threshold} word threshold")
        return ROUTE_TEXT, (
            f"the text layer yields {text_layer_words} words")

    # Nothing to go on, so default to the vision model. A missed diagram is
    # not recovered later, while an unnecessary call only costs time.
    return ROUTE_VISION, "no routing features available, defaulting to vision"


def summarise(report):
    """One plain line describing what a routed run did, for logs and the UI."""
    if not report:
        return ""
    total = report.get("pages", 0)
    vision = report.get("vision_pages", 0)
    if not total:
        return ""
    saved = total - vision
    return (f"{vision} of {total} slides sent to the vision model, "
            f"{saved} read from the text layer "
            f"({(saved / total):.0%} of model calls avoided).")


def extract_routed(input_path, client, dpi=None, max_pages=None,
                   picture_threshold=None, words_threshold=None,
                   on_page=None):
    """Extract a deck, choosing the path slide by slide.

    A PowerPoint file is converted to PDF once, so every input is handled the
    same way after that. Picture areas are read from the original pptx,
    because they cannot be recovered from a rendered page.

    A routed slide is rendered and sent to the vision model with the same
    retry and empty reply checks a whole-deck run uses. Slides not sent to the
    vision model are read from the PDF text layer without a model call.

    on_page(number, route, reason, seconds) is called after each slide, so a
    long run can report progress. Returns (text, report).
    """
    import time
    import tempfile
    import base64

    from modules import slide_vision, slides_pptx

    try:
        import fitz
    except ImportError as error:
        raise RuntimeError(
            "PyMuPDF is not installed. Install it with 'pip install pymupdf' "
            "to route slides page by page.") from error

    if dpi is None:
        import config
        dpi = config.VISION_DPI

    source = Path(str(input_path))
    picture_ratios = pptx_picture_area_ratios(source)

    workspace = None
    pdf_path = source
    try:
        if source.suffix.lower() == slides_pptx.PPTX_SUFFIX:
            workspace = tempfile.TemporaryDirectory()
            pdf_path = Path(slides_pptx.convert_to_pdf(str(source),
                                                       workspace.name))

        document = fitz.open(str(pdf_path))
        try:
            total = len(document)
            limit = total if not max_pages else min(total, max_pages)

            parts = []
            routes = []
            vision_pages = 0
            unreadable = []
            truncated = []
            retries = 0
            vision_seconds = 0.0

            for index in range(limit):
                number = index + 1
                page = document[index]
                layer_text = page.get_text("text") or ""
                ratio = picture_ratios.get(number)
                route, reason = decide(
                    picture_area_ratio=ratio,
                    text_layer_words=len(layer_text.split()),
                    picture_threshold=picture_threshold,
                    words_threshold=words_threshold)

                started = time.perf_counter()
                if route == ROUTE_VISION:
                    pixmap = page.get_pixmap(dpi=dpi)
                    url = ("data:image/png;base64," + base64.b64encode(
                        pixmap.tobytes("png")).decode("ascii"))
                    text, info = slide_vision.extract_one_image(url, client)
                    vision_pages += 1
                    retries += info["retries"]
                    if info["truncated"]:
                        truncated.append(number)
                    if info["unreadable"]:
                        unreadable.append(number)
                else:
                    text = layer_text.strip()
                elapsed = time.perf_counter() - started
                if route == ROUTE_VISION:
                    vision_seconds += elapsed

                routes.append({"page": number, "route": route,
                               "reason": reason,
                               "picture_area_ratio": ratio,
                               "text_layer_words": len(layer_text.split()),
                               "seconds": round(elapsed, 2)})
                header = (f"{slide_vision.PAGE_MARKER} {number} --- "
                          + ROUTE_MARKER.format(route=route))
                parts.append(f"{header}\n{text}".strip())

                if on_page is not None:
                    on_page(number, route, reason, elapsed)
        finally:
            document.close()
    finally:
        if workspace:
            workspace.cleanup()

    report = {
        "pages": limit,
        "vision_pages": vision_pages,
        "text_pages": limit - vision_pages,
        "unreadable_pages": unreadable,
        "truncated_pages": truncated,
        "retries": retries,
        "vision_seconds": round(vision_seconds, 2),
        "routes": routes,
    }
    return "\n\n".join(parts).strip(), report
