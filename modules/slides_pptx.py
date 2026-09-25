"""Read PowerPoint files, as text and as rendered pages.

Most lecture slides arrive as pptx, so the system accepts the deck the student
already has rather than asking for a PDF export first.

The module has two separate jobs, because one is cheap and one is not.

1. Text extraction through python-pptx. It installs with pip alone and reads
   the file format directly, in the same way pypdf reads a PDF. Slide
   boundaries, titles and table cells survive, where OCR flattens a rendered
   page into loose lines.

2. Rendering pages to images for the OCR and vision paths, which needs
   LibreOffice, since no pure Python library renders PowerPoint faithfully.
   This part is optional. Without LibreOffice the text path still works and the
   interface says what is missing, the same pattern as modules/audio_stt.py.

Speaker notes are read only when asked for and are off by default. No other
extraction path can see them, so including them would make the four-way
comparison unfair. It is the same reason each page is rendered once and shared
in evaluation/compare_extractors.py.
"""
import os
import shutil
import subprocess
import tempfile
from pathlib import Path

from modules.slide_vision import PAGE_MARKER

PPTX_SUFFIX = ".pptx"

# The default Windows install location, checked as well as PATH because the
# LibreOffice installer does not add itself to PATH there.
WINDOWS_SOFFICE = r"C:\Program Files\LibreOffice\program\soffice.exe"

# Seconds to allow for one conversion before giving up, so a stalled
# LibreOffice does not hang the interface.
CONVERT_TIMEOUT = 180

NOTES_MARKER = "[Speaker notes]"


def is_available():
    """Return whether pptx text extraction is possible, meaning python-pptx
    is installed.

    Rendering for the OCR and vision paths is reported separately by
    libreoffice_available, since the two have different dependencies.
    """
    try:
        import pptx  # noqa: F401  (python-pptx)
        return True
    except ImportError:
        return False


def find_soffice():
    """Locate the LibreOffice binary, or return None.

    Checks PATH first, then the default Windows install location, because the
    Windows installer does not put soffice.exe on PATH.
    """
    found = shutil.which("soffice")
    if found:
        return found
    if os.path.isfile(WINDOWS_SOFFICE):
        return WINDOWS_SOFFICE
    return None


def libreoffice_available():
    """Return whether pptx can be rendered to pixels (LibreOffice present)."""
    return find_soffice() is not None


def _shape_text(shape):
    """Return the text of one shape, including table cells.

    Tables are read cell by cell and joined with tabs, so the rows survive in
    the extracted text. That structure is what this path keeps and OCR loses.
    """
    parts = []
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append("\t".join(cells))
    elif getattr(shape, "has_text_frame", False):
        text = shape.text_frame.text.strip()
        if text:
            parts.append(text)
    return parts


def _slide_title(slide):
    try:
        placeholder = slide.shapes.title
    except (AttributeError, ValueError):
        return ""
    if placeholder is None:
        return ""
    return (placeholder.text or "").strip()


def _notes_text(slide):
    if not getattr(slide, "has_notes_slide", False):
        return ""
    frame = getattr(slide.notes_slide, "notes_text_frame", None)
    if frame is None:
        return ""
    return (frame.text or "").strip()


def extract_pptx_text(path, include_notes=False):
    """Extract slide text from a pptx file, keeping the slide boundaries.

    Uses the same page marker as modules/slide_vision.py, so output from the
    different extraction paths lines up when they are compared. The slide title
    comes first where there is one, then the remaining shape and table text in
    document order.

    include_notes is off by default, for the reason in the module docstring.
    When it is on, the notes are added under a labelled marker so they are not
    mistaken for text that was visible on the slide.

    Raises RuntimeError with an install hint when python-pptx is missing.
    """
    try:
        from pptx import Presentation
    except ImportError as error:
        raise RuntimeError(
            "python-pptx is not installed. Install it with "
            "'pip install python-pptx' to read PowerPoint files."
        ) from error
    presentation = Presentation(str(path))
    pages = []
    for number, slide in enumerate(presentation.slides, start=1):
        lines = []
        title = _slide_title(slide)
        if title:
            lines.append(title)
        for shape in slide.shapes:
            for text in _shape_text(shape):
                # The title placeholder is a shape too, so skip it here.
                if text != title:
                    lines.append(text)
        if include_notes:
            notes = _notes_text(slide)
            if notes:
                lines.append(f"{NOTES_MARKER}\n{notes}")
        body = "\n".join(lines).strip()
        pages.append(f"{PAGE_MARKER} {number} ---\n{body}".strip())
    return "\n\n".join(pages).strip()


def convert_to_pdf(path, output_dir):
    """Convert a pptx to PDF with LibreOffice and return the new path.

    Converting first means the existing PDF rendering path is reused, so the
    OCR and vision paths need to know nothing about PowerPoint. Raises
    RuntimeError with a clear message when LibreOffice is missing or the
    conversion fails, which is what the interface shows the user.
    """
    soffice = find_soffice()
    if soffice is None:
        raise RuntimeError(
            "LibreOffice was not found, so this PowerPoint file cannot be "
            "rendered to images for the OCR or vision paths. Either install "
            "LibreOffice, or export the deck to PDF from PowerPoint and "
            "upload that PDF instead. Native pptx text extraction still works "
            "without LibreOffice."
        )
    try:
        completed = subprocess.run(
            [soffice, "--headless", "--convert-to", "pdf", "--outdir",
             str(output_dir), str(path)],
            capture_output=True, text=True, timeout=CONVERT_TIMEOUT,
            check=False)
    except (OSError, subprocess.SubprocessError) as error:
        raise RuntimeError(
            f"LibreOffice conversion failed to start: {error}") from error
    produced = Path(output_dir) / (Path(str(path)).stem + ".pdf")
    if not produced.is_file():
        raise RuntimeError(
            "LibreOffice did not produce a PDF for this file "
            f"(exit code {completed.returncode}). "
            "Export the deck to PDF from PowerPoint and upload that instead."
        )
    return str(produced)


def render_pptx_to_images(path, dpi=None, max_pages=None):
    """Render pptx slides to images by converting to PDF first.

    The temporary PDF is deleted once the pages are rendered and only the
    images are returned, so the caller does not see the conversion.
    """
    from modules import slide_vision
    kwargs = {}
    if dpi is not None:
        kwargs["dpi"] = dpi
    if max_pages is not None:
        kwargs["max_pages"] = max_pages
    with tempfile.TemporaryDirectory() as workspace:
        pdf_path = convert_to_pdf(path, workspace)
        return slide_vision.render_pdf_to_images(pdf_path, **kwargs)
